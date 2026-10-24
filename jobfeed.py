"""Daily SDE-1 / Frontend job list for tier-1 Indian cities -> Excel.

Uses public company job-board APIs (Greenhouse, Lever, Ashby), which companies
publish for exactly this purpose. No login, no scraping.
Only jobs posted in the last 7 days are kept. Excel files go to daily_jobs/.
Usage:  python jobfeed.py            (new jobs since last run)
        python jobfeed.py --all      (every matching job from the last 7 days, ignore history)
"""
import json, re, sys, urllib.request, concurrent.futures as cf
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).parent
SEEN = HERE / "seen.json"
OUT_DIR = HERE / "daily_jobs"
MAX_AGE_DAYS = 7

ROLE = re.compile(r"\b(sde|software (development )?engineer|software developer|front[- ]?end|frontend|react|web developer|ui engineer)\b", re.I)
TOO_SENIOR = re.compile(r"\b(senior|sr\.?|lead|staff|principal|manager|head|director|architect|ii|iii|iv|sde[- ]?[23]|[23])\b|intern", re.I)
CITIES = {"Bengaluru": r"bengaluru|bangalore", "Hyderabad": r"hyderabad", "Pune": r"pune", "Mumbai": r"mumbai",
          "Delhi NCR": r"delhi|gurgaon|gurugram|noida", "Chennai": r"chennai", "Kolkata": r"kolkata", "Ahmedabad": r"ahmedabad",
          "Remote (India)": r"remote.*india|india.*remote"}

def city_of(loc):
    return next((c for c, p in CITIES.items() if re.search(p, loc or "", re.I)), None)

def wanted(title):
    return bool(ROLE.search(title)) and not TOO_SENIOR.search(title)

def recent(posted, today=None):
    """posted is an ISO date string (YYYY-MM-DD...); missing dates are dropped."""
    return bool(posted) and posted[:10] >= ((today or date.today()) - timedelta(days=MAX_AGE_DAYS)).isoformat()

def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "jobfeed/1.0"})
    with urllib.request.urlopen(req, timeout=20) as r:
        return json.load(r)

def fetch(line):
    ats, rest = line.split(":", 1)
    board, _, name = (p.strip() for p in rest.partition("|"))
    return [(name or board.title(),) + job[1:] for job in _fetch(ats, board)]

def _fetch(ats, board):
    if ats == "greenhouse":
        data = get(f"https://boards-api.greenhouse.io/v1/boards/{board}/jobs")["jobs"]
        return [(board, j["title"], j["location"]["name"], j["absolute_url"], (j.get("first_published") or j.get("updated_at") or "")[:10]) for j in data]
    if ats == "lever":
        data = get(f"https://api.lever.co/v0/postings/{board}?mode=json")
        return [(board, j["text"], j["categories"].get("location", ""), j["hostedUrl"],
                 date.fromtimestamp(j["createdAt"] / 1000).isoformat()) for j in data]
    if ats == "ashby":
        data = get(f"https://api.ashbyhq.com/posting-api/job-board/{board}")["jobs"]
        return [(board, j["title"], j.get("location", ""), j["jobUrl"], j.get("publishedAt", "")[:10]) for j in data]
    raise ValueError(f"unknown ATS '{ats}'")

def main():
    show_all = "--all" in sys.argv
    lines = [l.strip() for l in (HERE / "companies.txt").read_text().splitlines() if l.strip() and not l.startswith("#")]
    seen = set() if show_all or not SEEN.exists() else set(json.loads(SEEN.read_text()))
    rows, errors = [], []
    with cf.ThreadPoolExecutor(10) as ex:
        for line, fut in [(l, ex.submit(fetch, l)) for l in lines]:
            try:
                for company, title, loc, url, posted in fut.result():
                    city = city_of(loc)
                    if city and wanted(title) and recent(posted) and url not in seen:
                        rows.append((company, title, city, loc, posted, url))
            except Exception as e:
                errors.append(f"{line}: {e}")
    rows.sort(key=lambda r: r[4], reverse=True)
    for e in errors:
        print("  skipped", e)
    if not rows:
        return print(f"No new jobs from the last {MAX_AGE_DAYS} days.")

    wb = Workbook(); ws = wb.active; ws.title = "Jobs"
    ws.append(["Company", "Job Title", "City", "Location (as posted)", "Posted", "Apply Link"])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0A66C2")
    for r in rows:
        ws.append(r)
        cell = ws.cell(ws.max_row, 6); cell.hyperlink = r[5]; cell.font = Font(name="Arial", color="0563C1", underline="single")
    for col, w in zip("ABCDEF", [16, 50, 16, 34, 15, 60]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:F{max(ws.max_row, 2)}"
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f"jobs_{date.today().isoformat()}_{len(list(OUT_DIR.glob(f'jobs_{date.today()}*')))+1}.xlsx"
    wb.save(out)

    if not show_all:
        SEEN.write_text(json.dumps(sorted(seen | {r[5] for r in rows})))
    print(f"{len(rows)} {'matching' if show_all else 'new'} jobs -> daily_jobs/{out.name}")

if __name__ == "__main__":
    main()
